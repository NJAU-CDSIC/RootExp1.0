#!/usr/bin/env python3
"""
Benchmark Table 3: average GPU inference time per image (ms) at fixed 224×224.

Protocol:
  - batch size = 1
  - input = (1, 3, 224, 224) random tensor (architecture cost; matches training npy size)
  - warmup + timed repeats with torch.cuda.synchronize()
  - exclude disk I/O / CPU preprocess
  - OCRNet* = Yuan et al. OCRNet (HRNetV2-W48 + OCR); see tools/ocrnet_hrnet_w48.py

Usage (from repo root):
  /home/young/anaconda3/envs/vmunet/bin/python tools/benchmark_table3_inference_ms.py
  /home/young/anaconda3/envs/vmunet/bin/python tools/benchmark_table3_inference_ms.py --fill-docx
"""
from __future__ import annotations

import argparse
import csv
import importlib
import json
import os
import re
import subprocess
import sys
import time
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
TESTALL = REPO / "test_all_dataset"
ROOTBIN = REPO / "model" / "RootBin_Block_Mamba_EAM_ASPP"
DOCX_PATH = REPO / "Tables 3-new.docx"
CSV_PATH = REPO / "tools" / "table3_inference_ms.csv"
JSON_PATH = REPO / "tools" / "table3_inference_ms.json"

# Table column order → dataset folder
DATASETS = [
    ("Synthetic", "sumilte"),
    ("Wheat", "wheat"),
    ("PRMI", "prim"),
    ("RPT", "RPT"),
    ("BnaRoot", "da"),
]

# Table row order → (display name, folder under test_all_dataset/<ds>/ or special)
MODELS = [
    ("U-Net", "unet"),
    ("DeepLabv3", "Deeplabv3"),
    ("U-Net++", "unet++"),
    ("SegFormer", "SegFormer"),
    ("U-Mamba", "U-Mamba_adapted"),
    ("SegRoot", "SegRoot_adapted"),
    ("ITErRoot", "ITERroot_adapted"),
    ("OCRNet*", "OCRNet*"),  # HRNetV2-W48 + OCR (paper)
    ("ESegFormer", "ESegformer_adapted"),
    ("RootExp", "RootExp"),  # special
]

PYTHON = sys.executable


def pick_best_pth(model_dir: Path) -> Path | None:
    bests = sorted(model_dir.glob("best-epoch*-*.pth"))
    if not bests:
        p = model_dir / "best.pth"
        return p if p.is_file() else None

    def epoch_key(p: Path):
        m = re.search(r"best-epoch(\d+)", p.name)
        return int(m.group(1)) if m else -1

    return max(bests, key=epoch_key)


def detect_kind(name: str) -> str:
    n = name.lower()
    if name == "unet++" or n == "unet++":
        return "unet++"
    if name == "unet":
        return "unet"
    if n == "deeplabv3":
        return "deeplabv3"
    if "esegformer" in n:
        return "esegformer"
    if n == "segformer":
        return "segformer"
    if "u-mamba" in n or "umamba" in n:
        return "umamba"
    if "iterroot" in n:
        return "iterroot"
    if "segroot" in n:
        return "segroot"
    raise SystemExit(f"cannot map folder name to model kind: {name!r}")


def _purge_local_modules(model_root: Path):
    """Drop previously imported project modules so the next model_root wins."""
    root_s = str(model_root.resolve())
    doomed = []
    for name, mod in list(sys.modules.items()):
        path = getattr(mod, "__file__", None)
        if path and Path(path).resolve().as_posix().startswith(Path(root_s).as_posix()):
            doomed.append(name)
        elif name in {
            "configs",
            "configs.config_setting",
            "models",
            "utils",
            "engine",
            "loader",
        } or name.startswith("configs.") or name.startswith("models."):
            doomed.append(name)
    for name in doomed:
        sys.modules.pop(name, None)
    # Prefer this model_root at front of path
    while root_s in sys.path:
        sys.path.remove(root_s)
    sys.path.insert(0, root_s)


def _import_strip_predict():
    import importlib.util

    path = REPO / "tools" / "strip_predict_testall_one_model.py"
    spec = importlib.util.spec_from_file_location("strip_predict_testall_one_model", path)
    mod = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(mod)
    return mod


def build_baseline(model_root: Path, kind: str, device):
    # build_model inserts model_root itself; ensure clean import context first
    _purge_local_modules(model_root)
    mod = _import_strip_predict()
    return mod.build_model(model_root, kind, device)


def build_rootexp(device, lite: bool = False):
    _purge_local_modules(ROOTBIN)
    sys.path.insert(0, str(ROOTBIN))
    if lite:
        from configs.config_setting_lite import setting_config  # noqa: E402
    else:
        from configs.config_setting import setting_config  # noqa: E402
    from models.rootbin import RootBinBlockMambaEAMASPPNet  # noqa: E402

    cfg = setting_config.model_config
    m = RootBinBlockMambaEAMASPPNet(
        num_classes=cfg["num_classes"],
        input_channels=cfg["input_channels"],
        stage_name=cfg.get("stage_name", "stage0_unet_scblock_hvss_eam_aspp"),
        split_att=cfg.get("split_att", "fc"),
        encoder_channels=cfg.get("encoder_channels"),
        use_eam=cfg.get("use_eam", True),
        use_aspp=cfg.get("use_aspp", True),
        use_hvss=cfg.get("use_hvss", True),
        use_bridge=cfg.get("use_bridge", True),
        hvss_depth=cfg.get("hvss_depth", 1),
        drop_path_rate=cfg.get("drop_path_rate", 0.1),
        hvss_order=cfg.get("hvss_order", 2),
        hvss_d_state=cfg.get("hvss_d_state", 6),
        mlp_ratio=cfg.get("mlp_ratio", 1.5),
        use_skip_enhance=cfg.get("use_skip_enhance", True),
        skip_fusion_dim=cfg.get("skip_fusion_dim", 64),
        aspp_dilations=tuple(cfg.get("aspp_dilations", (1, 6, 12, 18))),
    )
    m = m.to(device)
    m.eval()
    return m


def load_weights(model, weight: Path):
    import torch

    ckpt = torch.load(str(weight), map_location="cpu")
    state = ckpt["model_state_dict"] if isinstance(ckpt, dict) and "model_state_dict" in ckpt else ckpt
    if isinstance(state, dict) and state and all(k.startswith("module.") for k in state):
        state = {k[len("module.") :]: v for k, v in state.items()}
    missing, unexpected = model.load_state_dict(state, strict=False)
    return len(missing), len(unexpected)


def measure_ms(model, device, hw=224, warmup=20, repeats=100) -> dict:
    import torch

    model_ref = model.module if hasattr(model, "module") else model
    model_ref.eval()
    dummy = torch.randn(1, 3, hw, hw, device=device)
    with torch.no_grad():
        for _ in range(warmup):
            _ = model_ref(dummy)
    torch.cuda.synchronize()
    t0 = time.perf_counter()
    with torch.no_grad():
        for _ in range(repeats):
            _ = model_ref(dummy)
    torch.cuda.synchronize()
    elapsed = time.perf_counter() - t0
    fps = repeats / elapsed
    ms = 1000.0 / fps
    return {"ms": ms, "fps": fps, "elapsed": elapsed, "repeats": repeats, "hw": hw}


def worker_main(args):
    import torch

    assert torch.cuda.is_available(), "CUDA required"
    device = torch.device("cuda:0")
    torch.backends.cudnn.benchmark = True

    if args.kind in ("rootexp", "rootexp_lite"):
        lite = args.kind == "rootexp_lite" or args.lite
        model = build_rootexp(device, lite=lite)
        weight = Path(args.weight) if args.weight else None
        # Full-width ckpt is incompatible with lite; only load when shapes match.
        if weight is not None and weight.is_file() and not lite:
            load_weights(model, weight)
        elif (not lite) and (ROOTBIN / "best-epoch1927-loss0.1049.pth").is_file() and not args.weight:
            load_weights(model, ROOTBIN / "best-epoch1927-loss0.1049.pth")
    elif args.kind == "ocrnet":
        here = Path(__file__).resolve().parent
        for d in (here, REPO / "tools", Path(__file__).resolve().parents[3] / "tools"):
            if (d / "ocrnet_hrnet_w48.py").is_file():
                sys.path.insert(0, str(d))
                break
        from ocrnet_hrnet_w48 import build_ocrnet

        model = build_ocrnet(num_classes=2, in_chans=3).to(device).eval()
    else:
        model_root = Path(args.model_root)
        model = build_baseline(model_root, args.kind, device)
        if args.weight:
            load_weights(model, Path(args.weight))

    stats = measure_ms(model, device, hw=args.hw, warmup=args.warmup, repeats=args.repeats)
    print(json.dumps(stats, ensure_ascii=False))


def run_worker(kind: str, model_root: Path | None, weight: Path | None, hw: int, warmup: int, repeats: int) -> dict:
    cmd = [
        PYTHON,
        str(Path(__file__).resolve()),
        "--worker",
        "--kind",
        kind,
        "--hw",
        str(hw),
        "--warmup",
        str(warmup),
        "--repeats",
        str(repeats),
    ]
    if model_root is not None:
        cmd += ["--model_root", str(model_root)]
    if weight is not None:
        cmd += ["--weight", str(weight)]
    env = os.environ.copy()
    env["CUDA_VISIBLE_DEVICES"] = env.get("CUDA_VISIBLE_DEVICES", "0")
    # Ensure repo root on PYTHONPATH for `from tools...` inside worker when needed
    env["PYTHONPATH"] = str(REPO) + (os.pathsep + env["PYTHONPATH"] if env.get("PYTHONPATH") else "")
    proc = subprocess.run(cmd, capture_output=True, text=True, cwd=str(REPO), env=env)
    if proc.returncode != 0:
        raise RuntimeError(
            f"worker failed kind={kind} root={model_root}\n"
            f"stdout:\n{proc.stdout[-2000:]}\nstderr:\n{proc.stderr[-4000:]}"
        )
    lines = [ln for ln in proc.stdout.strip().splitlines() if ln.strip().startswith("{")]
    if not lines:
        raise RuntimeError(f"no JSON from worker; stdout=\n{proc.stdout}\nstderr=\n{proc.stderr}")
    return json.loads(lines[-1])


def format_ms(ms: float) -> str:
    if ms >= 100:
        return f"{ms:.1f}"
    if ms >= 10:
        return f"{ms:.2f}"
    return f"{ms:.2f}"


def fill_docx(table_ms: dict, note: str):
    from docx import Document

    doc = Document(str(DOCX_PATH))
    if not doc.tables:
        raise SystemExit("no table in docx")
    table = doc.tables[0]
    # header already present; fill rows 1..
    name_to_row = {}
    for ri, row in enumerate(table.rows):
        label = row.cells[0].text.strip()
        if label:
            name_to_row[label] = ri

    col_names = [c[0] for c in DATASETS]
    for model_name, _ in MODELS:
        if model_name not in name_to_row:
            print(f"[warn] model row missing in docx: {model_name}")
            continue
        ri = name_to_row[model_name]
        for ci, col in enumerate(col_names, start=1):
            ms = table_ms[model_name][col]
            table.rows[ri].cells[ci].text = format_ms(ms)

    # append / update caption note as a paragraph after table if not present
    note_marker = "[Table 3 note]"
    found = False
    for p in doc.paragraphs:
        if note_marker in p.text:
            p.text = f"{note_marker} {note}"
            found = True
            break
    if not found:
        doc.add_paragraph(f"{note_marker} {note}")

    out = DOCX_PATH
    doc.save(str(out))
    print(f"wrote {out}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--worker", action="store_true")
    ap.add_argument("--kind", type=str, default="")
    ap.add_argument("--model_root", type=str, default="")
    ap.add_argument("--weight", type=str, default="")
    ap.add_argument("--hw", type=int, default=224)
    ap.add_argument("--warmup", type=int, default=20)
    ap.add_argument("--repeats", type=int, default=100)
    ap.add_argument("--fill-docx", action="store_true")
    ap.add_argument("--lite", action="store_true", help="RootExp-Lite narrower width")
    ap.add_argument(
        "--rootexp-variant",
        type=str,
        default="full",
        choices=["full", "lite"],
        help="Which RootExp width to put in Table 3",
    )
    ap.add_argument("--datasets", type=str, default="", help="comma subset of sumilte,wheat,prim,RPT,da")
    ap.add_argument("--only-rootexp", action="store_true", help="Only (re)measure RootExp row")
    args = ap.parse_args()

    if args.worker:
        worker_main(args)
        return

    ds_filter = None
    if args.datasets.strip():
        ds_filter = {x.strip() for x in args.datasets.split(",") if x.strip()}

    results = {}  # model -> col -> ms
    meta = {
        "device": None,
        "hw": args.hw,
        "warmup": args.warmup,
        "repeats": args.repeats,
        "protocol": "GPU forward only, batch=1, input 224x224, CUDA synchronize",
        "ocrnet_star": "HRNetV2-W48 + OCR head (paper assembly; tools/ocrnet_hrnet_w48.py)",
        "rootexp_variant": "lite" if (args.rootexp_variant == "lite" or args.lite) else "full",
    }
    if args.only_rootexp and JSON_PATH.is_file():
        prev = json.loads(JSON_PATH.read_text(encoding="utf-8"))
        results = prev.get("results_ms", {})
        meta = {**prev.get("meta", {}), **meta}

    # probe device name once
    probe = subprocess.run(
        [PYTHON, "-c", "import torch; print(torch.cuda.get_device_name(0))"],
        capture_output=True,
        text=True,
        cwd=str(REPO),
    )
    meta["device"] = probe.stdout.strip() if probe.returncode == 0 else "unknown"

    for col, ds in DATASETS:
        if ds_filter and ds not in ds_filter:
            continue
        print(f"\n===== dataset {col} ({ds}) =====")
        unet_ms = None
        iter_ms = None
        for model_name, folder in MODELS:
            results.setdefault(model_name, {})
            try:
                if args.only_rootexp and folder != "RootExp":
                    continue

                if folder == "OCRNet*":
                    stats = run_worker("ocrnet", None, None, args.hw, args.warmup, args.repeats)
                    ms = float(stats["ms"])
                    fps = 1000.0 / ms
                    print(f"  {model_name}: {ms:.3f} ms  (HRNetV2-W48+OCR)  ~{fps:.1f} FPS")
                    results[model_name][col] = ms
                    continue

                if folder == "RootExp":
                    use_lite = args.rootexp_variant == "lite" or args.lite
                    kind = "rootexp_lite" if use_lite else "rootexp"
                    weight = None
                    if not use_lite:
                        weight = ROOTBIN / "best-epoch1927-loss0.1049.pth"
                        if ds == "da":
                            da_w = (
                                REPO
                                / "DA_strip_inference_portable"
                                / "weights"
                                / "da_job01_best-epoch322-loss0.1250.pth"
                            )
                            if da_w.is_file():
                                weight = da_w
                    stats = run_worker(
                        kind,
                        None,
                        weight if (weight is not None and weight.is_file()) else None,
                        args.hw,
                        args.warmup,
                        args.repeats,
                    )
                else:
                    model_root = TESTALL / ds / folder
                    if not model_root.is_dir():
                        raise FileNotFoundError(model_root)
                    weight = pick_best_pth(model_root)
                    kind = detect_kind(folder)
                    stats = run_worker(kind, model_root, weight, args.hw, args.warmup, args.repeats)

                ms = float(stats["ms"])
                wname = weight.name if weight else ("lite-arch" if folder == "RootExp" else "none")
                print(f"  {model_name}: {ms:.3f} ms  ({stats['fps']:.1f} FPS)  weight={wname}")
                results[model_name][col] = ms
                if folder == "unet":
                    unet_ms = ms
                if folder == "ITERroot_adapted":
                    iter_ms = ms
            except Exception as e:
                print(f"  {model_name}: FAILED — {e}")
                results[model_name][col] = float("nan")

    # write CSV
    cols = [c for c, d in DATASETS if (not ds_filter or d in ds_filter)]
    # when only-rootexp, keep full column set from previous file if present
    if args.only_rootexp and results:
        cols = [c for c, _ in DATASETS]
    with open(CSV_PATH, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Method"] + cols + ["Unit"])
        for model_name, _ in MODELS:
            row = [model_name]
            for c in cols:
                v = results.get(model_name, {}).get(c, float("nan"))
                row.append(format_ms(v) if v == v else "NA")
            row.append("ms/image")
            w.writerow(row)
    print(f"\nwrote {CSV_PATH}")

    payload = {"meta": meta, "results_ms": results}
    JSON_PATH.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"wrote {JSON_PATH}")

    variant = meta.get("rootexp_variant", "full")
    lite_note = (
        " RootExp uses Lite width encoder_channels=[36,72,144,288,384] (~13.8M params); "
        "retrain before pairing with accuracy tables."
        if variant == "lite"
        else ""
    )
    note = (
        f"Measured on {meta['device']}; batch=1; input {args.hw}×{args.hw}; "
        f"warmup={args.warmup}, repeats={args.repeats}; GPU forward only (no I/O). "
        f"OCRNet* = HRNetV2-W48 + OCR (Yuan et al.).{lite_note}"
    )
    if args.fill_docx:
        fill_docx(results, note)
    else:
        print("skip docx fill (pass --fill-docx to write Tables 3-new.docx)")
        print("note:", note)


if __name__ == "__main__":
    main()
